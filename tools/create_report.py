from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Отчет_практическая_работа_Заметки.docx"


def set_cell_fill(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    props = table._tbl.tblPr
    borders = props.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def add_code(document: Document, lines: list[str]) -> None:
    for line in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.35)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.15

for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13)):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

styles["Title"].paragraph_format.space_after = Pt(18)
styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(8)
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(6)

footer = section.footer
footer.paragraphs[0].text = "Практическая работа  •  Flutter"
footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
footer.paragraphs[0].runs[0].font.size = Pt(9)
footer.paragraphs[0].runs[0].font.color.rgb = RGBColor(90, 90, 90)
add_page_number(footer.add_paragraph())

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Практическая работа")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(10)
run = subtitle.add_run("Разработка Flutter приложения Заметки")
run.bold = True
run.font.size = Pt(17)

level = doc.add_paragraph()
level.alignment = WD_ALIGN_PARAGRAPH.CENTER
level.paragraph_format.space_after = Pt(28)
level.add_run("Вариант на оценку 5").italic = True

meta = doc.add_table(rows=4, cols=2)
meta.autofit = False
meta.columns[0].width = Inches(2.0)
meta.columns[1].width = Inches(4.4)
set_table_borders(meta)
for i, (label, value) in enumerate(
    (
        ("Выполнил(а)", "____________________________"),
        ("Группа", "____________________________"),
        ("Проверил(а)", "____________________________"),
        ("Год", "2026"),
    )
):
    meta.cell(i, 0).text = label
    meta.cell(i, 1).text = value
    set_cell_fill(meta.cell(i, 0), "DDEEFF")
    for cell in meta.rows[i].cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    meta.cell(i, 0).paragraphs[0].runs[0].bold = True

intro = doc.add_paragraph()
intro.paragraph_format.space_before = Pt(30)
intro.add_run(
    "В ходе работы создано простое мобильное приложение для ввода, просмотра, "
    "редактирования и удаления заметок. Интерфейс выполнен в голубой цветовой гамме."
)

doc.add_page_break()

doc.add_heading("1 Цель работы", level=1)
doc.add_paragraph(
    "Получить практические навыки создания интерфейса на Flutter, работы с состоянием StatefulWidget, "
    "коллекциями и диалоговыми окнами."
)

doc.add_heading("2 Постановка задачи", level=1)
for item in (
    "реализовать поле для ввода текста;",
    "добавить кнопку Сохранить;",
    "выводить сохраненные заметки в виде списка;",
    "добавить редактирование и удаление заметок;",
    "оформить приложение в голубой цветовой гамме.",
):
    add_bullet(doc, item)

doc.add_heading("3 Средства разработки", level=1)
tools = doc.add_table(rows=1, cols=2)
tools.autofit = False
tools.columns[0].width = Inches(2.0)
tools.columns[1].width = Inches(4.4)
headers = ("Средство", "Назначение")
for index, header in enumerate(headers):
    tools.cell(0, index).text = header
    set_cell_fill(tools.cell(0, index), "B8DDF7")
    tools.cell(0, index).paragraphs[0].runs[0].bold = True
for tool, purpose in (
    ("Flutter", "Создание кроссплатформенного приложения"),
    ("Dart", "Язык программирования"),
    ("Material 3", "Компоненты интерфейса и цветовая тема"),
    ("flutter_test", "Автоматическая проверка поведения интерфейса"),
):
    cells = tools.add_row().cells
    cells[0].text = tool
    cells[1].text = purpose
for row_index, row in enumerate(tools.rows):
    for cell in row.cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if row_index > 0 and row_index % 2 == 0:
        for cell in row.cells:
            set_cell_fill(cell, "F2F8FF")
set_table_borders(tools)

doc.add_heading("4 Реализация", level=1)
doc.add_heading("4.1 Структура проекта", level=2)
doc.add_paragraph(
    "Точка входа и весь учебный интерфейс находятся в файле lib/main.dart. В классе NotesApp заданы тема и главный экран. "
    "Класс NotesPage хранит список заметок и обновляет экран через setState."
)

doc.add_heading("4.2 Добавление заметки", level=2)
doc.add_paragraph(
    "Метод _saveNote удаляет лишние пробелы, проверяет ввод и добавляет текст в список. Пустое значение не сохраняется."
)
add_code(
    doc,
    [
        "final String text = _noteController.text.trim();",
        "if (text.isEmpty) { ... return; }",
        "setState(() {",
        "  _notes.add(text);",
        "  _noteController.clear();",
        "});",
    ],
)

doc.add_heading("4.3 Редактирование и удаление", level=2)
doc.add_paragraph(
    "Кнопка с карандашом открывает диалог с текущим текстом. После подтверждения элемент списка заменяется. "
    "Кнопка с корзиной показывает окно подтверждения и удаляет заметку только при выборе кнопки Удалить."
)

doc.add_heading("4.4 Интерфейс", level=2)
doc.add_paragraph(
    "Для оформления использована тема Material 3 с голубым основным цветом. Поле ввода, кнопка и карточки заметок имеют скругленные углы. "
    "Счетчик рядом с заголовком Заметки показывает число элементов."
)

doc.add_heading("5 Тестирование", level=1)
doc.add_paragraph(
    "Проект проверен статическим анализатором и двумя виджет-тестами. По результатам flutter analyze ошибок не найдено. "
    "Команда flutter test завершилась успешно: 2 теста пройдены. Дополнительно выполнена веб-сборка командой flutter build web; компиляция прошла успешно."
)
tests = doc.add_table(rows=1, cols=3)
tests.autofit = False
tests.columns[0].width = Inches(0.7)
tests.columns[1].width = Inches(4.8)
tests.columns[2].width = Inches(1.0)
for index, header in enumerate(("№", "Проверка", "Результат")):
    tests.cell(0, index).text = header
    set_cell_fill(tests.cell(0, index), "B8DDF7")
    tests.cell(0, index).paragraphs[0].runs[0].bold = True
for number, check in (
    ("1", "Добавление, редактирование и удаление заметки"),
    ("2", "Запрет сохранения пустой заметки"),
):
    cells = tests.add_row().cells
    cells[0].text = number
    cells[1].text = check
    cells[2].text = "Пройден"
for row_index, row in enumerate(tests.rows):
    for cell_index, cell in enumerate(row.cells):
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if cell_index in (0, 2):
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if row_index == 2:
        for cell in row.cells:
            set_cell_fill(cell, "F2F8FF")
set_table_borders(tests)

doc.add_heading("6 Запуск проекта", level=1)
doc.add_paragraph("В корневой папке проекта нужно выполнить команды:")
add_code(doc, ["flutter pub get", "flutter run"])
doc.add_paragraph("Для повторной проверки кода:")
add_code(doc, ["flutter analyze", "flutter test"])

doc.add_heading("7 Вывод", level=1)
doc.add_paragraph(
    "Все требования задания выполнены. Пользователь может создавать заметки, видеть их в списке, редактировать и удалять. "
    "Приложение имеет единый голубой дизайн, а код проходит анализ и автотесты."
)

doc.add_heading("8 Состав прилагаемых файлов", level=1)
for item in (
    "папка Flutter-проекта с исходным кодом;",
    "файл README.md с краткой инструкцией;",
    "файл test/widget_test.dart с автотестами;",
    "данный отчет.",
):
    add_bullet(doc, item)

doc.save(OUTPUT)
print(OUTPUT)
